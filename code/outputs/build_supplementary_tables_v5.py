import argparse, copy, glob, os, re
import numpy as np, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment

ap = argparse.ArgumentParser()
ap.add_argument("--style", default="data/tables/Supplementary_Tables_v4_final.xlsx")
ap.add_argument("--out", default="data/tables/Supplementary_Tables_v5.xlsx")
ap.add_argument("--select-csv", default="data/results/v3/select_main.csv")
A = ap.parse_args()
r1 = lambda v: np.round(v, 1)
r2 = lambda v: np.round(v, 2)
NR = "Not reached"


def style_table(n):

    x = pd.read_excel(A.style, f"Supplementary Table {n}", header=1)
    return x.iloc[:-1].reset_index(drop=True)


t5 = style_table(5)
NAME = dict(zip(t5.ISO3, t5.Country))
NAME["GLB"] = "Global"
CR = pd.read_csv("data/results/v3/mc/summary_crossover_ratios.csv").set_index("iso3")
ORDER = ["GLB"] + CR.drop("GLB").sort_values("xext_median_%").index.tolist()


def order(df, col="ISO3"):
    df = df.copy()
    df["_o"] = df[col].map({k: i for i, k in enumerate(ORDER)})
    return df.sort_values("_o", kind="stable").drop(columns="_o").reset_index(drop=True)


N_DRAWS = len(pd.read_csv("data/results/v3/mc/draws.csv"))


def cx(v):

    return NR if (pd.isna(v) or not np.isfinite(v) or v > 100) else round(float(v), 1)


NEW = {}
LEGEND = {}
HEADER = {}

t3 = style_table(3)
cal = pd.read_csv("data/results/v3/decomp/main_check_calibration.csv")
err4 = (
    cal[["err_inc_I", "err_inc_S", "err_d_I", "err_d_S", "err_d_all"]].abs().max().max()
)
full = pd.read_csv("data/results/v3/decomp/main_check_calibration_full.csv")
ev = 100 * pd.concat([full.err_d_I, full.err_d_S]).abs()
CAL = dict(median=float(ev.median()), max=float(ev.max()))
rows = [
    [
        "Calibration, four-state life table",
        "All countries, sexes and five-year age groups",
        "Largest relative difference from GBD 2023 incidence, cause-specific and all-cause death rates",
        f"{err4:.1e}",
        "Numerical tolerance 1e-7",
    ],
    [
        "Calibration, full model with each state divided by diabetes status",
        "All countries, sexes and five-year age groups",
        "Relative difference in first-year IHD and IS deaths from the calibrated rates (%), median (largest)",
        f"{CAL['median']:.2f} ({CAL['max']:.2f})",
        "Largest differences in the oldest age groups",
    ],
]
cc = pd.read_csv("data/results/v3/grid/coverage_check_summary.csv")
cc["alloc"] = cc.run.str.split("_").str[0]
cov_rows = {}
for al, g in cc.groupby("alloc"):
    lab = {"risk": "Risk-prioritised allocation", "unif": "Uniform allocation"}.get(al)
    if lab:
        cov_rows[lab] = [
            round(float(g.max_excess_before_pp.max()), 3),
            f"{int((g.years_overshoot > 0).sum())} of {len(g)} run-country combinations",
        ]
t3 = t3.copy()
for i, r in t3.iterrows():
    if r.Check == "Annual coverage check" and r["Group or country"] in cov_rows:
        t3.loc[i, ["Model value", "Comparison"]] = cov_rows[r["Group or country"]]
NEW[3] = pd.concat([pd.DataFrame(rows, columns=t3.columns), t3], ignore_index=True)
LEGEND[3] = [
    (
        "External validation simulated",
        "Calibration compares modelled first-year rates with GBD 2023 targets in every "
        "country, sex and five-year age group, for the four-state life table and for the full model with each state "
        "divided by diabetes status. External validation simulated",
    ),
    (
        "Calibration error was below 1e-7 and population conservation error below 1e-15 in every country, sex and age (not tabulated).",
        "Population conservation error was below 1e-15 in every country, sex and age (not tabulated).",
    ),
]

t4 = style_table(4)
dr = pd.read_csv("data/results/v3/mc/draws.csv")
KEYS = [
    "hr_mi",
    "hr_stroke",
    "hr_cvd",
    "hr_mi_dm",
    "hr_stroke_dm",
    "hr_cvd_dm",
    "hr_t2d",
    "s_ihd",
    "stop_mult",
    "restart_mult",
    "resid_cv",
    "resid_t2d",
    "dm_hr_ihd",
    "dm_hr_is",
    "sbp_shift",
    "sbp_shift_dm",
]
assert len(t4) == len(KEYS), "template Table 4 rows do not match the parameter list"
rows = [
    [
        t4.Parameter[i],
        t4.Distribution[i],
        r2(dr[k].median()),
        r2(dr[k].quantile(0.025)),
        r2(dr[k].quantile(0.975)),
    ]
    for i, k in enumerate(KEYS)
]
u = pd.read_csv("data/params/diet_paf_uncertainty.csv").set_index("outcome")
for oc, lab in [
    ("IHD", "IHD"),
    ("IschStroke", "ischaemic stroke"),
    ("T2D", "type 2 diabetes"),
    ("CKD", "CKD"),
]:
    k = f"diet_factor_{oc}"
    if k in dr:
        rows.append(
            [
                f"Dietary effect multiplier, {lab}",
                f"Split log-normal, median 1, 95% range {u.lower[oc] / u.val[oc]:.2f} to {u.upper[oc] / u.val[oc]:.2f}",
                r2(dr[k].median()),
                r2(dr[k].quantile(0.025)),
                r2(dr[k].quantile(0.975)),
            ]
        )
NEW[4] = pd.DataFrame(rows, columns=t4.columns)
LEGEND[4] = [
    ("Each of 500 simulations", f"Each of {N_DRAWS} simulations"),
    (
        "Uncertainty in dietary effects was propagated through the potential impact fractions.",
        "Dietary effects were multiplied by an outcome-specific factor, fitted to the GBD 2023 95% uncertainty interval of "
        "the diet-attributable fraction relative to its point estimate, common to all countries within each draw and "
        "drawn independently for each outcome.",
    ),
    ("CV, cardiovascular;", "CKD, chronic kidney disease; CV, cardiovascular;"),
]

t7 = style_table(7)
pdl = pd.read_csv("data/results/v3/supp/prob_diet_leads.csv", dtype={"cov": str}).pivot(
    index="iso3", columns="cov", values="p_diet_leads"
)
ov = pd.read_csv("data/results/v3/mc_traj/overtaking_year.csv").pivot(
    index="iso3", columns="run", values="first_year_diet_ahead"
)
rows = []
for iso, r in CR.iterrows():
    d = [
        NAME.get(iso, iso),
        iso,
        cx(r["xext_median_%"]),
        cx(r["xext_lo_%"]),
        cx(r["xext_hi_%"]),
        round(100 * r.share_not_reached),
        round(100 * r["P(crossover<20%)"]),
        round(r.treat100_pct_of_diet_median),
        r2(r.r20_median),
        r2(r.r20_lo),
        r2(r.r20_hi),
    ]
    d += [round(100 * pdl.loc[iso, c]) for c in ("0.100", "0.200", "0.500", "afford")]
    d += [
        "Not within 30 years" if pd.isna(ov.loc[iso, c]) else int(ov.loc[iso, c])
        for c in ("c0.10", "c0.20", "cafford")
    ]
    rows.append(d)
NEW[7] = order(pd.DataFrame(rows, columns=t7.columns))
HEADER[7] = {
    "Simulations with crossover above 100% (%)": "Simulations in which crossover was not reached (%)"
}
LEGEND[7] = [
    (
        "Where treatment did not match dietary policy even with every eligible adult enrolled, the value above 100% is the "
        "size of programme needed relative to enrolling all eligible adults, calculated as the deaths averted by dietary "
        "policy divided by those averted at full coverage. This is a lower bound, because any additional people treated "
        "would be at lower risk.",
        "Where treatment did not match dietary policy even with every eligible adult enrolled, crossover is reported as not "
        "reached, and treatment at full coverage gives the deaths averted by treatment with every eligible adult enrolled "
        "as a percentage of those averted by dietary policy.",
    ),
    ("share of 500 simulations", f"share of {N_DRAWS} simulations"),
]

t8 = style_table(8)
ob = pd.read_csv("data/results/v3/supp/outcomes_by_arm.csv", dtype={"cov": str})
OUTC = [
    ("deaths", "Deaths from IHD and ischaemic stroke"),
    ("t2d", "Incident type 2 diabetes cases"),
    ("ckd", "Deaths from CKD"),
]
ARM = [
    ("diet", "Dietary policy (feasible package)"),
    ("glp1", "GLP-1RA treatment"),
    ("combo", "Both (independent action)"),
]
COVL = [("0.100", "10%"), ("0.200", "20%"), ("afford", "Maximum affordable")]
obi = ob.set_index(["iso3", "cov", "arm", "outcome"])
rows = []
for iso in ORDER:
    for oc, ol in OUTC:
        for arm, al in ARM:
            for cv, cl in ([("0.200", "Not applicable")] if arm == "diet" else COVL):
                q = obi.loc[(iso, cv, arm, oc)]
                rows.append(
                    [
                        NAME.get(iso, iso),
                        iso,
                        ol,
                        al,
                        cl,
                        r2(q["median"] / 1e3),
                        r2(q.lo / 1e3),
                        r2(q.hi / 1e3),
                    ]
                )
NEW[8] = pd.DataFrame(rows, columns=t8.columns)
LEGEND[8] = [("from 500 simulations", f"from {N_DRAWS} simulations")]

t9 = style_table(9)
o = pd.read_csv("data/results/v3/overlap/overlap_by_country.csv").assign(
    outcome="IHD+IS"
)
o2 = pd.read_csv("data/results/v3/mc_traj_t2dckd/overlap_by_outcome.csv")
o2 = o2[o2.outcome != "IHD+IS"]
OL = {
    "IHD+IS": "Deaths from IHD and ischaemic stroke",
    "T2D": "Incident type 2 diabetes cases",
    "CKD": "Deaths from CKD",
}
rows = []
for oc in ("IHD+IS", "T2D", "CKD"):
    src_ = o if oc == "IHD+IS" else o2[o2.outcome == oc]
    for _, r in src_.iterrows():
        has_ui = "add_diet_to_trt_pct_lo" in r and pd.notna(
            r.get("add_diet_to_trt_pct_lo", np.nan)
        )
        rows.append(
            [
                NAME.get(r.iso3, r.iso3),
                r.iso3,
                OL[oc],
                r.coverage,
                float(r.overlap_pct),
                r2(r.combined_median / 1e3),
                r2(r.combined_lo / 1e3),
                r2(r.combined_hi / 1e3),
                r1(r.add_diet_to_trt_pct_median),
                r2(1 + r.add_diet_to_trt_pct_median / 100),
                r2(1 + r.add_diet_to_trt_pct_lo / 100) if has_ui else "Not estimated",
                r2(1 + r.add_diet_to_trt_pct_hi / 100) if has_ui else "Not estimated",
            ]
        )
x9 = pd.DataFrame(rows, columns=t9.columns)
x9["_c"] = x9[t9.columns[3]].map({"10%": 0, "20%": 1, "Maximum affordable": 2})
x9["_oc"] = x9.Outcome.map({v: i for i, v in enumerate(OL.values())})
NEW[9] = order(x9.sort_values(["_oc", "_c", t9.columns[4]]).drop(columns=["_c", "_oc"]))
LEGEND[9] = [("from 500 simulations", f"from {N_DRAWS} simulations")]

t10 = style_table(10)
s = pd.read_csv("data/results/v3/mc_t2dckd/summary_outcome_crossover.csv")
rows = []
for _, r in s.iterrows():
    ne = r.outcome == "CKD" and r.r20_median <= 0
    rows.append(
        [
            NAME.get(r.iso3, r.iso3),
            r.iso3,
            {"T2D": "Incident type 2 diabetes", "CKD": "Deaths from CKD"}[r.outcome],
        ]
        + (
            ["Not estimable"] * 3
            if ne
            else [cx(r.xext_median_pct), cx(r.xext_lo_pct), cx(r.xext_hi_pct)]
        )
        + (["Not estimable"] if ne else [round(100 * r.share_not_reached)])
        + (
            ["Not estimable"] * 3
            if ne
            else [r2(r.r20_median), r2(r.r20_lo), r2(r.r20_hi)]
        )
    )
x10 = pd.DataFrame(rows, columns=t10.columns)
x10["_oc"] = x10.Outcome.map({"Incident type 2 diabetes": 0, "Deaths from CKD": 1})
NEW[10] = order(x10.sort_values("_oc", kind="stable").drop(columns="_oc"))
HEADER[10] = {
    "Simulations with crossover above 100% (%)": "Simulations in which crossover was not reached (%)"
}
LEGEND[10] = [
    (
        "averted over 30 years.",
        "averted over 30 years. Where treatment did not match dietary policy even with every "
        "eligible adult enrolled, crossover is reported as not reached.",
    )
]

t11 = style_table(11)
f_ = pd.read_csv("data/tables/table_single_factor_v3.csv")
code2name = dict(zip(f_["GDD variable code"], f_["Dietary factor"]))
x11 = f_.drop(columns=["GDD variable code"]).rename(columns={"Country": "ISO3"})
x11.insert(0, "Country", x11.ISO3.map(NAME))
assert len(x11.columns) == len(t11.columns)
x11.columns = t11.columns
fo = {v: i for i, v in enumerate(pd.unique(t11["Dietary factor"]))}
NEW[11] = order(
    x11.assign(_f=x11["Dietary factor"].map(fo))
    .sort_values("_f", kind="stable")
    .drop(columns="_f")
)
t12 = style_table(12)
rows = []
for f in sorted(glob.glob("data/results/v3/single_factor/single_v[0-9][0-9].csv")):
    x = pd.read_csv(f)
    x = x[(x.arm == "diet") & x.outcome.isin(["IHD", "IschStroke"])]
    yr = [c for c in x.columns if c.startswith("yr")]
    g = x.groupby("iso3")[yr].sum()
    fac = code2name.get(re.search(r"single_(v\d+)", f).group(1))
    for iso, v in g.iterrows():
        cum = np.cumsum(v.values)
        rows.append(
            [NAME.get(iso, iso), iso, fac]
            + [round(float(cum[k - 1])) for k in (5, 10, 15, 20, 25, 30)]
        )
x12 = pd.DataFrame(rows, columns=t12.columns)
NEW[12] = order(
    x12.assign(_f=x12["Dietary factor"].map(fo))
    .sort_values("_f", kind="stable")
    .drop(columns="_f")
)

sc = pd.read_csv("data/results/v3/mc_scen/summary_scenarios.csv")
d3 = pd.read_csv("data/results/v3/mc_scen/summary_scenarios_disc3.csv")
sc = pd.concat(
    [sc, d3[d3.scenario == "main"].assign(scenario="discount_3pct")], ignore_index=True
)
FM = {
    "cvd": "as for cardiovascular death",
    "same": "as for non-fatal events",
    "none": "none",
}
G = {
    "Main analysis": {"main": "Main analysis"},
    "Treatment persistence": {
        "const_retention": "Fixed retention of two thirds of the effect",
        "resid_0": "No benefit retained after stopping",
        "no_restart": "No restarting after stopping",
        "restart_late0": "No restarting after two years off",
        "persist_full": "Nobody stops treatment",
        "restart_late0.5": "Restarting halved after two years off",
        "stop_x1.5": "Stopping 1.5 times as often",
        "stop_x0.5": "Stopping half as often",
    },
    "Effects after disease onset": {
        "post_A": "No effect on deaths after onset",
        "post_C": "Nor on fatal first events",
        "diet_post_na": "Dietary sodium effect also on deaths after onset",
        "diet_post_na_wg": "Dietary sodium and whole grain effects also on deaths after onset",
    },
    "Dietary policy implementation": {
        "phase_12": "Phase-in over 12 years",
        "phase_16": "Phase-in over 16 years",
        "maintain_0.4": "Policy effect held at 40%",
        "maintain_0.6": "Policy effect held at 60%",
        "maintain_0.8": "Policy effect held at 80%",
    },
    "Scope of dietary policy": {
        "aspirational": "Aspirational scenario",
        "tailored": "Package tailored to each country",
        "sodium30": "Sodium reduction of 30%",
        "tailored_cap": "Tailored package, with benefit in people with IHD or IS capped at US level",
    },
    "Eligible population and programme": {
        "tier3_bmi27": "All adults with BMI 27 or more eligible",
        "bmi25": "Eligibility extended to BMI 25 or more",
        "bmi30": "Eligibility restricted to BMI 30 or more",
        "baseline_cohort": "No replacement of enrolled adults who die",
        "no_dm_tier": "No priority for diabetes",
        "select_like": "Only adults resembling SELECT",
        "uniform": "Places allocated uniformly",
        "dynamic": "Every eligible adult with IHD or IS enrolled from year 2",
        "cap": "Absolute benefit in people with IHD or IS capped at US level",
    },
    "Trial effect estimates": {
        "primary_weak": "40% of the effect without CVD",
        "class_overall": "Class-level estimate for everyone",
        "itt_select": "Correction for discontinuation in trials",
        "sattar_dm": "Earlier meta-analysis, with diabetes",
        "tier3_none": "No effect without CVD or diabetes",
        "tier2_rewind": "REWIND, diabetes without CVD",
        "select_only": "SELECT only, without diabetes",
    },
    "Model structure": {
        "is_cf10": "Case fatality of first ischaemic stroke 10%",
        "is_cf20": "Case fatality of first ischaemic stroke 20%",
        "dm_no_risk": "Diabetes does not raise cardiovascular risk",
        "dm_pr2": "Diabetes prevalence ratio 2",
        "dm_pr3": "Diabetes prevalence ratio 3",
        "k0.8": "Relative duration of IHD or IS with BMI ≥27, 0.8",
        "k1.2": "Relative duration of IHD or IS with BMI ≥27, 1.2",
        **{
            f"s{a_}_{m_}": f"IHD deaths as fatal first events {int(round(float(a_) * 100))}%, effect on fatal first events {FM[m_]}"
            for a_ in ("0.25", "0.4", "0.55")
            for m_ in FM
        },
        "bmi_gamma": "Gamma distribution of BMI",
        "or2": "Odds ratio of having both IHD and IS, 2",
        "or4": "Odds ratio of having both IHD and IS, 4",
        "rr1.5": "Risk of the second condition after the first, 1.5",
        "rr2": "Risk of the second condition after the first, 2",
    },
    "Population and time": {
        "discount_3pct": "Discounting 3% a year",
        "open_pop": "Open population (new adults enter)",
    },
}
LBL = {k: (g, v) for g, d_ in G.items() for k, v in d_.items()}


def scen_rows(df):
    out = []
    for _, r in df.iterrows():
        share = (
            [r1(100 * r.t20_median), r1(100 * r.t20_lo), r1(100 * r.t20_hi)]
            if pd.notna(r.get("t20_median"))
            else [np.nan] * 3
        )
        if pd.notna(r.r20_median) and 0 < r.r20_median < 0.05:          # dietary benefit near zero: share unstable
            share = ["Not estimable"] * 3
        cross = (
            ["Not estimated (20% coverage only)"] * 3
            if r.scenario == "dynamic"
            else [cx(100 * r.xext_median), cx(100 * r.xext_lo), cx(100 * r.xext_hi)]
        )
        ratio = (                                                       # treatment benefit near or below zero
            ["Not estimable"] * 3
            if r.scenario == "select_like" or (pd.notna(r.r20_lo) and r.r20_lo <= 0)
            else [r2(r.r20_median), r2(r.r20_lo), r2(r.r20_hi)]
        )
        row = cross + ratio + [round(100 * r.P_diet_leads_20)] + share
        if pd.notna(r.r20_median) and r.r20_median == 0:
            row = (
                ["No modelled dietary benefit"] * 6
                + [row[6]]
                + ["No modelled dietary benefit"] * 3
            )
        out.append(row)
    return out


t13 = style_table(13)
s13 = sc[~sc.scenario.str.startswith("sw_")].copy()
miss = sorted(set(s13.scenario) - set(LBL))
assert not miss, f"scenarios without a label: {miss}"
x13 = pd.DataFrame(
    [
        [LBL[k][0], LBL[k][1], NAME.get(i, i), i] + v
        for k, i, v in zip(s13.scenario, s13.iso3, scen_rows(s13))
    ],
    columns=t13.columns,
)
GORD = {g: i for i, g in enumerate(G)}
x13 = order(x13)
x13["_g"] = x13["Scenario group"].map(GORD)
NEW[13] = (
    x13.sort_values(["_g", "Scenario"], kind="stable")
    .drop(columns="_g")
    .reset_index(drop=True)
)
LEGEND[13] = [
    (
        "Crossover coverage above 100% is the programme size needed relative to enrolling all eligible adults (lower bound).",
        "Where treatment did not match dietary policy even with every eligible adult enrolled, crossover is reported as not reached.",
    )
]
t14 = style_table(14)
sw = sc[sc.scenario.str.startswith("sw_")].copy()
m = sw.scenario.str.extract(r"sw_pm([\d.]+)_stop([\d.]+)_rd(\d+)")
x14 = pd.DataFrame(
    [
        [int(round(float(a) * 100)), float(b), int(c), NAME.get(i, i), i] + v
        for a, b, c, i, v in zip(m[0], m[1], m[2], sw.iso3, scen_rows(sw))
    ],
    columns=t14.columns,
)
x14 = x14.sort_values(
    list(t14.columns[:3]), ascending=[False, True, True], kind="stable"
)
x14["_o"] = x14.ISO3.map({k: i for i, k in enumerate(ORDER)})
NEW[14] = (
    x14.sort_values(
        list(t14.columns[:3]) + ["_o"],
        ascending=[False, True, True, True],
        kind="stable",
    )
    .drop(columns="_o")
    .reset_index(drop=True)
)
LEGEND[14] = [
    (
        "Medians and 95% uncertainty intervals from 300 simulations per scenario.",
        "Medians and 95% uncertainty intervals from 300 simulations per scenario. Where treatment did not match dietary "
        "policy even with every eligible adult enrolled, crossover is reported as not reached.",
    )
]

E = "data/results/v3/equity_cost_final/cov0.20"
t19 = style_table(19)
e = pd.read_csv(f"{E}/tableS_equity_v3.csv")
e.insert(0, "Country", e.ISO3.map(NAME))
e["Diet-to-treatment ratio"] = (
    e["Diet-to-treatment ratio"]
    .astype(object)
    .where(e["Diet-to-treatment ratio"].notna(), "Treatment not affordable")
)
assert list(e.columns) == list(t19.columns), (list(e.columns), list(t19.columns))
NEW[19] = order(e)
t20 = style_table(20)
x20 = pd.read_csv(f"{E}/tableS_equity_mapping_sensitivity_v3.csv")
assert list(x20.columns) == list(t20.columns)
NEW[20] = x20
t21 = style_table(21)
CV = {"0.10": "10%", "0.20": "20%", "afford": "Maximum affordable"}
c_ = pd.concat(
    pd.read_csv(f"data/results/v3/equity_cost_final/cov{c}/tableS_cost_v3.csv").assign(
        **{"Coverage of eligible adults": l}
    )
    for c, l in CV.items()
)
dk, tk = (
    "Dietary policy expenditure per death averted (USD)",
    "GLP-1RA expenditure per death averted (USD)",
)
ref = c_[c_["Coverage of eligible adults"] == "20%"].set_index(
    ["ISO3", "Price scenario", "Outcomes"]
)[dk]
c_[dk] = [
    ref.get((i, p, o_), v)
    for i, p, o_, v in zip(c_.ISO3, c_["Price scenario"], c_.Outcomes, c_[dk])
]
c_["Ratio, treatment to diet"] = (
    pd.to_numeric(c_[tk], errors="coerce") / pd.to_numeric(c_[dk], errors="coerce")
).round(1)
c_.insert(0, "Country", c_.ISO3.map(NAME))
c_ = c_[["Country", "ISO3"] + [c for c in c_.columns if c not in ("Country", "ISO3")]]
assert list(c_.columns) == list(t21.columns), (list(c_.columns), list(t21.columns))
NEW[21] = order(c_.reset_index(drop=True))
t22 = style_table(22)
q_ = pd.read_csv(
    "data/results/v3/sa_t2d_ckd/summary_sa_t2d_ckd.csv", dtype={"coverage": str}
)
q_ = q_[q_.coverage.isin(["0.20", "0.2", "afford"])].dropna(subset=["t2d_glp1"])
SL = {
    "main": "Main analysis",
    "t2d_bucur026": "GLP-1RA network meta-analysis estimate",
    "t2d_dual010": "Dual GIP/GLP-1 receptor agonist estimate",
    "sbp_select331": "Blood pressure fall from SELECT",
    "sbp_step510": "Blood pressure fall from STEP 1",
    "sbp_meta632": "Larger blood pressure fall",
    "sbp_dm0": "No blood pressure fall with diabetes",
    "sbp_none": "No blood pressure fall",
}
lab22 = t22.drop_duplicates("Analysis").set_index("Analysis")[
    [t22.columns[1], t22.columns[2]]
]
rows = []
for k, l in SL.items():
    for cov, cl in (("0.20", "20%"), ("afford", "Maximum affordable")):
        r = q_[
            (q_.scenario == k)
            & (q_.coverage.isin([cov, "0.2"] if cov == "0.20" else [cov]))
        ]
        if not len(r):
            continue
        r = r.iloc[0]
        rows.append(
            [
                l,
                lab22.loc[l].iloc[0],
                lab22.loc[l].iloc[1],
                cl,
                round(r.t2d_glp1 / 1e3),
                round(r.t2d_diet / 1e3),
                r2(r.t2d_diet / r.t2d_glp1) if r.t2d_glp1 > 0 else "Not estimable",
                r1(r.ckd_glp1 / 1e3),
                r1(r.ckd_diet / 1e3),
                (
                    r2(r.ckd_diet / r.ckd_glp1)
                    if r.ckd_glp1 > 0
                    else "No treatment effect on CKD"
                ),
            ]
        )
NEW[22] = pd.DataFrame(rows, columns=t22.columns)

SORT_KEYS = {
    8: ["Outcome", "Arm", "Coverage of eligible adults"],
    9: ["Outcome", "Coverage of eligible adults", "Overlap of effects (%)"],
    10: ["Outcome"],
    11: ["Dietary factor"],
    12: ["Dietary factor"],
    13: ["Scenario group", "Scenario"],
    14: [
        "Dietary policy effect maintained (%)",
        "Stopping probability multiplier",
        "Dietary policy phase-in (years)",
    ],
    19: ["Price scenario", "Education group"],
    21: ["Price scenario", "Outcomes", "Coverage of eligible adults"],
}


def like_template(n, df):

    t = style_table(n)
    if n == 13:                                                       # group renamed since the template
        t["Scenario group"] = t["Scenario group"].replace({"Effect in people with IHD or IS": "Effects after disease onset"})
    if "ISO3" not in df or "ISO3" not in t:
        return df
    keys = SORT_KEYS.get(n, [])
    df = df.copy()
    df["_c"] = df.ISO3.map({k: i for i, k in enumerate(ORDER)})
    if not keys:
        return (
            df.sort_values("_c", kind="stable")
            .drop(columns="_c")
            .reset_index(drop=True)
        )
    norm = (
        lambda x: x[keys]
        .apply(
            lambda c: c.map(
                lambda v: (
                    f"{float(v):g}"
                    if isinstance(v, (int, float, np.number))
                    else str(v)
                )
            )
        )
        .agg("|".join, axis=1)
    )
    rank = {k: i for i, k in enumerate(pd.unique(norm(t)))}
    df["_k"] = norm(df).map(rank)
    if n == 13 and df["_k"].isna().any():                            # scenarios added since the template: end of their group
        last = {}
        for k_, v_ in rank.items():
            last[k_.split("|")[0]] = max(last.get(k_.split("|")[0], -1), v_)
        miss_ = df["_k"].isna(); lab_ = list(pd.unique(df.loc[miss_, "Scenario"]))
        df.loc[miss_, "_k"] = [last[g] + 0.1 * (1 + lab_.index(x)) for g, x in zip(df.loc[miss_, "Scenario group"], df.loc[miss_, "Scenario"])]
    assert df["_k"].notna().all(), (
        n,
        "rows not found in the template",
        list(norm(df)[df["_k"].isna()][:3]),
    )
    country_major = t.ISO3.iloc[0] == t.ISO3.iloc[1]
    return (
        df.sort_values(["_c", "_k"] if country_major else ["_k", "_c"], kind="stable")
        .drop(columns=["_k", "_c"])
        .reset_index(drop=True)
    )


for n in list(NEW):
    if n not in (3, 4, 20, 22):
        NEW[n] = like_template(n, NEW[n])

wb = load_workbook(A.style)
report = []
for n, df in sorted(NEW.items()):
    ws = wb[f"Supplementary Table {n}"]
    tmpl = style_table(n)
    hdr = [ws.cell(2, j).value for j in range(1, ws.max_column + 1)]
    nc = len([h for h in hdr if h is not None])
    assert nc == len(df.columns), (n, nc, len(df.columns))
    assert [str(h) for h in hdr[:nc]] == [str(c) for c in df.columns], (
        n,
        "header mismatch",
    )
    last = ws.max_row
    legend = ws.cell(last, 1).value or ""
    for a_, b_ in LEGEND.get(n, []):
        assert a_ in legend, (n, "legend text not found", a_[:60])
        legend = legend.replace(a_, b_)
    data_style = [copy.copy(ws.cell(3, j)._style) for j in range(1, nc + 1)]
    leg_style, leg_h = (
        copy.copy(ws.cell(last, 1)._style),
        ws.row_dimensions[last].height,
    )
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row > 1:
            ws.unmerge_cells(str(rng))
    ws.delete_rows(3, last - 2)
    for j, h in enumerate(df.columns, 1):
        ws.cell(2, j).value = HEADER.get(n, {}).get(h, h)
    for i, row in enumerate(df.itertuples(index=False), 3):
        for j, v in enumerate(row, 1):
            v = (
                None
                if (isinstance(v, float) and np.isnan(v))
                else (v.item() if hasattr(v, "item") else v)
            )
            c = ws.cell(i, j, v)
            c._style = copy.copy(data_style[j - 1])
    lr = len(df) + 3
    c = ws.cell(lr, 1, legend)
    c._style = leg_style
    ws.merge_cells(start_row=lr, start_column=1, end_row=lr, end_column=nc)
    ws.row_dimensions[lr].height = (
        leg_h if leg_h else 15 * max(2, len(legend) // 150 + 1)
    )
    report.append((n, len(tmpl), len(df)))
DROP = {1}
RENUM = {2: 1, 3: 2, 4: 3, 5: 4, 6: 5, 8: 6, 7: 7}
RENUM.update({o: o - 1 for o in range(9, 23)})
assert sorted(RENUM.values()) == list(range(1, 22)) and not (set(RENUM) & DROP)
for o in DROP:
    wb.remove(wb[f"Supplementary Table {o}"])
for o in RENUM:
    wb[f"Supplementary Table {o}"].title = f"tmp_{o}"
for o, n in RENUM.items():
    ws = wb[f"tmp_{o}"]
    ws.title = f"Supplementary Table {n}"
    ws.cell(1, 1).value = re.sub(
        r"^Supplementary Table \d+",
        f"Supplementary Table {n}",
        str(ws.cell(1, 1).value),
    )
cs = wb["Contents"]
titles = {}
for r in range(2, cs.max_row + 1):
    v = cs.cell(r, 1).value
    m_ = re.search(r"\d+", str(v)) if v else None
    if m_ and int(m_.group()) in RENUM:
        titles[RENUM[int(m_.group())]] = cs.cell(r, 2).value
last_row = cs.max_row
for i, n in enumerate(sorted(titles), 2):
    cs.cell(i, 1).value = f"Supplementary Table {n}"
    cs.cell(i, 2).value = titles[n]
for r in range(len(titles) + 2, last_row + 1):
    for c in range(1, cs.max_column + 1):
        cs.cell(r, c).value = None
wb._sheets = [cs] + sorted(
    [w for w in wb.worksheets if w.title != "Contents"],
    key=lambda w: int(w.title.split()[-1]),
)
# ---- how each table's estimates were obtained (labels and notes added after renumbering) ----------------------
from openpyxl.styles import Alignment


def add_note(n, text=None, at="end", replace=None):
    """Add a sentence to a table's note: at the start, before the abbreviation list starting with `at`, or at the end."""
    ws = wb[f"Supplementary Table {n}"]
    c = ws.cell(ws.max_row, 1)
    v = str(c.value or "")
    if replace:
        assert replace[0] in v, (n, "note text not found", replace[0][:60])
        v = v.replace(*replace)
    if text and at == "start":
        v = text + " " + v
    elif text and at != "end":
        assert v.count(at) == 1, (n, "abbreviation list not found", at)
        v = v.replace(at, " " + text + at)
    elif text:
        v = v + " " + text
    c.value = v
    nc = max(1, sum(1 for h in ws[2] if h.value is not None))
    ws.row_dimensions[ws.max_row].height = max(ws.row_dimensions[ws.max_row].height or 15,
                                                15 * (len(v) // max(60, 14 * nc) + 2))


def relabel(n, col, old, new):
    ws = wb[f"Supplementary Table {n}"]
    k = 0
    for r in range(3, ws.max_row):
        if ws.cell(r, col).value == old:
            ws.cell(r, col).value = new
            k += 1
    assert k, (n, "label not found", old)


DRAWS300 = ("The main-analysis estimates in this table are based on the first 300 of the 1000 main-analysis simulations, "
            "so that every scenario is compared with the same simulations. The primary estimates, based on all 1000 "
            "simulations, are reported in Supplementary Table 7 and differ slightly.")
relabel(12, 2, "Main analysis", "Main analysis (first 300 simulations)")
add_note(12, DRAWS300, at=" UI, uncertainty interval.")
add_note(13, "The combination of 100% dietary policy effect, stopping probability multiplier 1 and an 8-year phase-in "
             "corresponds to the main analysis in these 300 simulations. " + DRAWS300.split(". ", 1)[1],
         at=" UI, uncertainty interval.")
DET = "Deterministic projections with central parameter values."
for n in (10, 11):
    add_note(n, DET, at="start")
add_note(18, "Totals across the 19 countries exclude the rest of the world, so they are not comparable with the global "
             "estimates in Supplementary Table 6.", at=" GLP-1RA, glucagon-like")
relabel(21, 1, "Main analysis", "Main analysis (deterministic)")   # becomes Table 22
add_note(21, replace=("Deterministic analyses, global totals over 30 years.",
                    "Deterministic analyses with central parameter values, global totals over 30 years. The main-analysis "
                    "rows therefore differ slightly from the Monte Carlo medians in Supplementary Table 6."))
cs.cell(len(titles) + 3, 1).value = (
    "Three kinds of estimates are reported. Main results (Supplementary Tables 6 to 9) are medians and 95% uncertainty "
    "intervals from 1000 Monte Carlo simulations. Scenario and sensitivity analyses (Supplementary Tables 12 and 13) "
    "use the first 300 of these simulations, with the main analysis recomputed on the same 300 simulations for "
    "comparison. Single-factor, equity, expenditure and secondary-outcome sensitivity analyses (Supplementary Tables "
    "10, 11 and 18 to 21) are deterministic projections with central parameter values. Small differences between "
    "tables in the same quantity therefore reflect the number of simulations or the use of central values, not "
    "different versions of the model.")
r_ = len(titles) + 3
cs.merge_cells(start_row=r_, start_column=1, end_row=r_, end_column=2)
cs.cell(r_, 1).alignment = Alignment(wrap_text=True, vertical="top")
cs.row_dimensions[r_].height = 75

# ---- corrected education PAFs, equal-budget table and renumbering (October 2026) -------------------------------
import copy as _copy
EG = {"Low": 1, "Medium": 2, "High": 3}
_v2 = pd.read_csv("data/results/v3/equity_paf_by_edu_v2.csv")
PAF2 = {(r.outcome, r.iso3, int(r.edu)): float(r.paf) for r in _v2.itertuples()}
# Table 15: education-specific PAFs from the corrected dose-response curves
ws = wb["Supplementary Table 15"]; k_ = 0
for r in range(3, ws.max_row):
    oc, iso, g = ws.cell(r, 1).value, ws.cell(r, 3).value, ws.cell(r, 4).value
    if g in EG:
        ws.cell(r, 5).value = round(PAF2[(oc, iso, EG[g])], 2); k_ += 1
assert k_ == ws.max_row - 3, ("Table 15 rows not all updated", k_)
add_note(15, None, replace=("These are the PAF values from which the low-to-high education ratios in the preceding table are derived.",
                            "The low-to-high education ratios in Supplementary Table 14 are the ratios of these PAFs."))
# Table 14: central ratio of the Table 15 PAFs, intervals from 200 draws of the education gradient
_R = pd.read_csv("data/results/v3/equity_paf_ratio_ui.csv").set_index(["outcome", "iso3"])
_fmt = lambda v: "Not estimable" if (pd.isna(v) or not np.isfinite(v) or v > 100) else round(float(v), 3)
ws = wb["Supplementary Table 14"]; ws.cell(2, 4).value = "Low/high education PAF ratio"; k_ = 0
for r in range(3, ws.max_row):
    key = (ws.cell(r, 1).value, ws.cell(r, 3).value)
    if key in _R.index:
        x_ = _R.loc[key]; ws.cell(r, 4).value, ws.cell(r, 5).value, ws.cell(r, 6).value = _fmt(x_.central), _fmt(x_.lo), _fmt(x_.hi); k_ += 1
assert k_ == ws.max_row - 3, ("Table 14 rows not all updated", k_)
add_note(14, None, replace=("Ratios are reported as medians with 95% uncertainty intervals from 200 Monte Carlo draws over the GDD "
                            "exposure distributions, with dose-response curves held fixed.",
                            "Ratios are the ratios of the education-specific PAFs in Supplementary Table 15, with 95% uncertainty "
                            "intervals from 200 Monte Carlo draws over the GDD exposure distributions. Dose-response curves were held "
                            "fixed, so these intervals reflect uncertainty in the education gradient only and are not comparable with "
                            "the intervals of the main analysis. Upper bounds above 100, where the high-education PAF approaches zero, "
                            "are reported as not estimable."))
# Table 16: IHD rows recomputed with the corrected PAFs; population shares recovered from the unchanged stroke rows
ws = wb["Supplementary Table 16"]
rows16 = {(ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value): r for r in range(3, ws.max_row) if ws.cell(r, 3).value in EG}
share = {(iso, g): ws.cell(r, 6).value / ws.cell(r, 4).value for (oc, iso, g), r in rows16.items() if oc == "IschStroke" and ws.cell(r, 4).value}
def _recompute(oc, write):
    worst = 0.0
    for iso in sorted({k[1] for k in rows16 if k[0] == oc}):
        rs = {g: rows16[(oc, iso, g)] for g in EG if (oc, iso, g) in rows16}
        paf = {g: PAF2[(oc, iso, EG[g])] for g in rs}; w = {g: paf[g] * share[(iso, g)] for g in rs}; tot = sum(w.values())
        for g, r in rs.items():
            vals = [round(paf[g], 2), round(w[g], 2), round(w[g] / tot, 3), round(w[g] / tot * ws.cell(r, 8).value / 1000, 1)]
            if write: ws.cell(r, 4).value, ws.cell(r, 6).value, ws.cell(r, 7).value, ws.cell(r, 9).value = vals
            else: worst = max(worst, max(abs(a - b) for a, b in zip(vals, [ws.cell(r, c).value for c in (4, 6, 7, 9)])))
    return worst
assert _recompute("IschStroke", write=False) <= 0.0101, "the Table 16 method does not reproduce the stroke rows"
_recompute("IHD", write=True)
# Table 21 (new): equal-budget comparison; the former Table 21 becomes Table 22
old21 = wb["Supplementary Table 21"]; old21.title = "Supplementary Table 22"
old21.cell(1, 1).value = str(old21.cell(1, 1).value).replace("Supplementary Table 21.", "Supplementary Table 22.", 1)
bq = pd.read_csv("data/results/v3/expenditure_ui/budget_ui.csv")
PORD = {"Branded": 0, "Generic $140/yr": 1, "Generic $28/yr": 2}
bq = bq.assign(_c=bq.ISO3.map({k: i for i, k in enumerate(ORDER)}), _p=bq.price.map(PORD)).sort_values(["_c", "_p"])
UI3 = lambda h: [h, "95% UI lower", "95% UI upper"]
H21 = (["Country", "ISO3", "GLP-1RA price", "Dietary policy cost over 30 years (US$ millions)"] + UI3("Deaths averted by dietary policy")
       + ["A. Treatment coverage bought with the dietary policy cost (% of eligible adults)"] + UI3("A. Deaths averted by that treatment")
       + ["B. Budget, treatment at 20% coverage (US$ millions)", "B. Dietary policy cost as share of the budget (%)"]
       + UI3("B. Deaths averted, treatment only") + ["B. Treatment coverage with the remaining budget (% of eligible adults)"]
       + UI3("B. Deaths averted, dietary policy plus treatment") + UI3("C. Deaths averted per dollar, dietary policy relative to treatment"))
NA = "Not applicable (dietary policy cost exceeds the budget)"
def _n(v, d=0):
    return "Not estimable" if pd.isna(v) or not np.isfinite(v) else (round(float(v)) if d == 0 else round(float(v), d))
def _row21(r):
    cA = "Above 100" if pd.isna(r.A_coverage) else round(100 * r.A_coverage, 1)
    B_ = [NA] * 4 if r.diet_exceeds_budget else [round(100 * r.B_remaining_coverage, 1), _n(r.B_both), _n(r.B_both_lo), _n(r.B_both_hi)]
    A_ = [_n(r.A_trt), _n(r.A_trt_lo), _n(r.A_trt_hi)] if pd.notna(r.A_coverage) else ["Not estimable"] * 3
    return ([NAME.get(r.ISO3, r.ISO3), r.ISO3, r.price, round(r.diet_cost_musd, 1), _n(r.diet), _n(r.diet_lo), _n(r.diet_hi), cA] + A_
            + [round(r.B_budget_musd, 1), round(100 * r.B_diet_share, 1), _n(r.B_trt), _n(r.B_trt_lo), _n(r.B_trt_hi)] + B_
            + [_n(r.C, 2), _n(r.C_lo, 2), _n(r.C_hi, 2)])
ws21 = wb.create_sheet("Supplementary Table 21", index=wb.sheetnames.index("Supplementary Table 22"))
src = wb["Supplementary Table 20"]
def _style(dst, s_):
    dst.font, dst.fill, dst.alignment, dst.border = _copy.copy(s_.font), _copy.copy(s_.fill), _copy.copy(s_.alignment), _copy.copy(s_.border)
    dst.number_format = s_.number_format
STY = {k: src.cell(*rc) for k, rc in dict(title=(1, 1), head=(2, 1), data=(3, 1), note=(src.max_row, 1)).items()}
def write_table(ws, title, header, rows, note, widths=None, head_h=75):
    """Write a supplementary table sheet in the house style (title, header, rows, note)."""
    for rng in list(ws.merged_cells.ranges): ws.unmerge_cells(str(rng))
    if ws.max_row: ws.delete_rows(1, ws.max_row)
    ws.cell(1, 1, title); _style(ws.cell(1, 1), STY["title"]); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
    for jj, h in enumerate(header, 1): ws.cell(2, jj, h); _style(ws.cell(2, jj), STY["head"])
    ws.row_dimensions[2].height = head_h
    for ii, row in enumerate(rows, 3):
        for jj, v in enumerate(row, 1): ws.cell(ii, jj, v); _style(ws.cell(ii, jj), STY["data"])
    lr = len(rows) + 3; ws.cell(lr, 1, note); _style(ws.cell(lr, 1), STY["note"])
    ws.merge_cells(start_row=lr, start_column=1, end_row=lr, end_column=len(header))
    ws.row_dimensions[lr].height = max(60, 15 * (len(note) // (14 * len(header)) + 2))
    for jj, w in enumerate(widths or [], 1): ws.column_dimensions[ws.cell(2, jj).column_letter].width = w
T21 = "Supplementary Table 21. Equal-budget comparison of dietary policy and GLP-1RA treatment, by country and GLP-1RA price"
UINOTE = ("Medians and 95% uncertainty intervals from 1000 simulations. Expenditure is at central parameter values and deaths "
          "averted come from each simulation, so the intervals reflect uncertainty in the health effects, and ratios are "
          "computed within each simulation.")
leg = (UINOTE + " Expenditure was calculated as in Supplementary Table 20, treatment paying only for time on drug and dietary "
       "policy costing US$0.55 per adult per year, over 30 years without discounting. Deaths are deaths from IHD and ischaemic "
       "stroke averted over 30 years. A, the GLP-1RA coverage of eligible adults whose 30-year expenditure equals the cost of "
       "the dietary package, and the deaths it would avert. B, the expenditure of treatment at 20% coverage spent on treatment "
       "alone, or on the dietary package first with the remainder on treatment, using the combination of the two under "
       "independent action. C, deaths averted per dollar spent by dietary policy relative to treatment at 20% coverage, which equals "
       "the treatment-to-diet ratio of expenditure per death averted in Supplementary Table 20, and is also the factor by which "
       "the cost of the package could rise before treatment became cheaper per death averted. "
       "Coverage was not limited to the maximum affordable level, because the programme is assumed to be publicly financed, and "
       "deaths were interpolated between the coverage levels of the Monte Carlo analysis. GLP-1RA, glucagon-like peptide-1 "
       "receptor agonist; IHD, ischaemic heart disease; UI, uncertainty interval.")
write_table(ws21, T21, H21, [_row21(r) for r in bq.itertuples()], leg, widths=[22, 7, 14, 13] + [12] * 23)
# contents: the new table and the renumbered one, and the general note
r22 = next(r for r in range(2, cs.max_row + 1) if cs.cell(r, 1).value == "Supplementary Table 21")
cs.cell(r22 + 1, 1).value, cs.cell(r22 + 1, 2).value = "Supplementary Table 22", cs.cell(r22, 2).value
cs.cell(r22, 2).value = T21.split(". ", 1)[1]
for c in (1, 2): _style(cs.cell(r22 + 1, c), cs.cell(r22, c))
note_r = next(r for r in range(2, cs.max_row + 1) if str(cs.cell(r, 1).value or "").startswith("Three kinds of estimates"))
txt = ("Three kinds of estimates are reported. Main results (Supplementary Tables 6 to 9, 20 and 21) are medians and 95% "
       "uncertainty intervals from 1000 Monte Carlo simulations, with expenditure in Supplementary Tables 20 and 21 at central "
       "parameter values. Scenario, sensitivity and education analyses (Supplementary Tables 12, 13, 18 and 19) use the first "
       "300 of these simulations, with the main analysis recomputed on the same 300 simulations for comparison. Single-factor "
       "rankings and the secondary-outcome sensitivity analyses (Supplementary Tables 10, 11 and 22) are deterministic "
       "projections with central parameter values. Small differences between tables in the same quantity therefore reflect "
       "the number of simulations or the use of central values, not different versions of the model.")
for rng in list(cs.merged_cells.ranges):
    if rng.min_row == note_r: cs.unmerge_cells(str(rng))
cs.cell(note_r, 1).value = None; cs.cell(note_r + 1, 1).value = txt
cs.merge_cells(start_row=note_r + 1, start_column=1, end_row=note_r + 1, end_column=2)
cs.cell(note_r + 1, 1).alignment = Alignment(wrap_text=True, vertical="top"); cs.row_dimensions[note_r + 1].height = 80
# notes that pre-empt arithmetic on summarised values
add_note(8, "Within each simulation, the combination under independent action averted slightly fewer deaths than the sum of "
            "the two interventions alone, because they partly act on the same people. Its median can nonetheless exceed the "
            "sum of the separate medians, because medians are not additive.", at=" CKD, chronic kidney disease;")
add_note(12, "Where the treatment benefit was close to or below zero in some simulations, the diet-to-treatment ratio is not "
             "estimable and the treatment benefit as a percentage of the dietary benefit should be read instead, and where the "
             "dietary benefit was close to zero the percentage is not estimable.", at=" UI, uncertainty interval.")

add_note(5, "Negative single-factor estimates, which GBD 2023 gives for red meat and some stroke types, were counted as "
            "zero for each outcome.", at=" CV, cardiovascular;")
# ---- uncertainty intervals for the education, expenditure and sensitivity tables (Tables 18-20) ------------------
PRORD = {"Branded": 0, "Generic $140/yr": 1, "Generic $28/yr": 2}; EORD = {"Low education": 0, "Medium education": 1, "High education": 2}
def _note(n): ws_ = wb[f"Supplementary Table {n}"]; return str(ws_.cell(ws_.max_row, 1).value or "")
def _r(v, d):
    return "Not estimable" if pd.isna(v) or not np.isfinite(v) else round(float(v), d)
# Table 18: deaths averted by education group
eq = pd.read_csv("data/results/v3/equity_mc/equity_ui.csv")
eq = eq.assign(_c=eq.ISO3.map({k: i for i, k in enumerate(ORDER)}), _p=eq["Price scenario"].map(PRORD), _e=eq["Education group"].map(EORD)).sort_values(["_c", "_p", "_e"])
H18 = (["Country", "Price scenario", "ISO3", "Education group", "Population share (25+)", "PAF ratio", "Affordable fraction"]
       + UI3("Dietary policy averted (thousands)") + UI3("GLP-1RA averted if affordable to all (thousands)")
       + UI3("GLP-1RA averted (thousands)") + UI3("Diet-to-treatment ratio"))
def _row18(r):
    rat = (["Treatment not affordable"] * 3 if r.glp1_median == 0 else [_r(r.ratio_median, 2), _r(r.ratio_lo, 2), _r(r.ratio_hi, 2)])
    return ([NAME.get(r.ISO3, r.ISO3), r[0], r.ISO3, r[2], r[3], r[4], r[5]]
            + [_r(r.diet_median, 2), _r(r.diet_lo, 2), _r(r.diet_hi, 2), _r(r.glp1_all_median, 2), _r(r.glp1_all_lo, 2), _r(r.glp1_all_hi, 2),
               _r(r.glp1_median, 2), _r(r.glp1_lo, 2), _r(r.glp1_hi, 2)] + rat)
n18 = _note(18)
assert "Treatment at 20% coverage of eligible adults." in n18
n18 = n18.replace("Treatment at 20% coverage of eligible adults.",
                  "Treatment at 20% coverage of eligible adults. Medians and 95% uncertainty intervals from the first 300 simulations of "
                  "the main analysis, with education-specific PAFs, population shares, affordability and the education-income link at "
                  "central values, and ratios computed within each simulation.")
write_table(wb["Supplementary Table 18"], str(wb["Supplementary Table 18"].cell(1, 1).value), H18,
            [_row18(r) for r in eq.drop(columns=["_c", "_p", "_e"]).itertuples(index=False)], n18, widths=[18, 14, 7, 16, 11, 9, 11] + [11] * 12)
# Table 19: sensitivity to the education-income link
mp = pd.read_csv("data/results/v3/equity_mc/mapping_ui.csv")
H19 = (["Education-income link"] + UI3("Dietary policy, low education (thousands)") + UI3("GLP-1RA, low education (thousands)")
       + UI3("Diet-to-treatment ratio") + ["Countries with no treatment benefit"])
rows19 = [[r["Education-income link"], _r(r.diet_median, 0), _r(r.diet_lo, 0), _r(r.diet_hi, 0), _r(r.glp1_median, 0), _r(r.glp1_lo, 0),
           _r(r.glp1_hi, 0), _r(r.ratio_median, 2), _r(r.ratio_lo, 2), _r(r.ratio_hi, 2), r["Countries with no treatment benefit"]] for _, r in mp.iterrows()]
write_table(wb["Supplementary Table 19"], str(wb["Supplementary Table 19"].cell(1, 1).value), H19, rows19,
            "Medians and 95% uncertainty intervals from the first 300 simulations of the main analysis, with ratios computed within "
            "each simulation. " + _note(19), widths=[34] + [12] * 9 + [20])
# Table 20: expenditure per death averted
cu = pd.read_csv("data/results/v3/expenditure_ui/cost_ui.csv")
CORD = {"10%": 0, "20%": 1, "Maximum affordable": 2}; OORD = {"IHD + IS": 0, "IHD + IS + CKD": 1}
cu = cu.assign(_c=cu.ISO3.map({k: i for i, k in enumerate(ORDER)}), _p=cu.price.map(PRORD), _v=cu.coverage.map(CORD), _o=cu.outcomes.map(OORD)).sort_values(["_c", "_p", "_v", "_o"])
H20 = (["Country", "ISO3", "Price scenario", "Outcomes", "Coverage of eligible adults", "Discounting", "Enrolled person-years (millions)",
        "Paying person-years (millions)"] + UI3("GLP-1RA expenditure per death averted (USD)")
       + UI3("Dietary policy expenditure per death averted (USD)") + UI3("Ratio, treatment to diet"))
rows20 = [[NAME.get(r.ISO3, r.ISO3), r.ISO3, r.price, r.outcomes, r.coverage, "None", round(r.enrolled_py_m, 2), round(r.paying_py_m, 2),
           _r(r.trt_pd, 0), _r(r.trt_pd_lo, 0), _r(r.trt_pd_hi, 0), _r(r.diet_pd, 0), _r(r.diet_pd_lo, 0), _r(r.diet_pd_hi, 0),
           _r(r.ratio, 1), _r(r.ratio_lo, 1), _r(r.ratio_hi, 1)] for r in cu.itertuples()]
_fd = "data/results/v3/expenditure_ui/cost_ui_disc3.csv"
if os.path.exists(_fd):                                   # costs and deaths discounted at 3% a year (300 simulations)
    cd_ = pd.read_csv(_fd); py_ = {(r.ISO3, r.price, r.coverage): (round(r.enrolled_py_m, 2), round(r.paying_py_m, 2)) for r in cu[cu.outcomes == "IHD + IS"].itertuples()}
    pos = {(r[1], r[2], r[4], r[3]): i for i, r in enumerate(rows20)}
    for r in cd_.itertuples():
        e_, p_ = py_[(r.ISO3, r.price, r.coverage)]
        new = [NAME.get(r.ISO3, r.ISO3), r.ISO3, r.price, "IHD + IS", r.coverage, "3% a year", e_, p_, _r(r.trt_pd, 0), _r(r.trt_pd_lo, 0),
               _r(r.trt_pd_hi, 0), _r(r.diet_pd, 0), _r(r.diet_pd_lo, 0), _r(r.diet_pd_hi, 0), _r(r.ratio, 1), _r(r.ratio_lo, 1), _r(r.ratio_hi, 1)]
        k0 = pos.get((r.ISO3, r.price, r.coverage, "IHD + IS + CKD"))
        rows20.insert(k0 + 1 if k0 is not None else len(rows20), new); pos = {(x[1], x[2], x[4], x[3] if x[5] == "None" else "disc"): i for i, x in enumerate(rows20)}
    print(f"Table 20: {len(cd_)} discounted rows added")
n20 = ("Rows without discounting are medians and 95% uncertainty intervals from 1000 simulations, and rows with costs and deaths "
       "discounted at 3% a year (mid-year) from the first 300, as in the discounting scenario. " if os.path.exists(_fd) else "")
n20 = (n20 + (UINOTE.replace("Medians and 95% uncertainty intervals from 1000 simulations. ", "") if n20 else UINOTE) + " In one of the 1000 simulations, treatment averted no deaths in some countries because the sampled hazard ratios "
       "were above 1, and its expenditure per death averted was then counted as unbounded. " + _note(20))
write_table(wb["Supplementary Table 20"], str(wb["Supplementary Table 20"].cell(1, 1).value), H20, rows20, n20,
            widths=[18, 7, 14, 14, 12, 11, 11, 11] + [12] * 9)

# ---- review of the supplementary tables (October 2026): sources as footnote markers, notes, panels -------------
def _setnote(n, txt):
    ws_ = wb[f"Supplementary Table {n}"]; ws_.cell(ws_.max_row, 1).value = txt
def _rows(n):
    ws_ = wb[f"Supplementary Table {n}"]; return ws_, [[c.value for c in r] for r in ws_.iter_rows(min_row=3, max_row=ws_.max_row - 1)]
# Table 1: sources as footnote markers, full references in the note
SRC1 = {"WHO guideline 2025; Galli et al, JACC 2025": "WHO guideline†; meta-analysis*",
        "Product labelling; SELECT": "Product labelling‡; SELECT trial§",
        "Product labelling": "Product labelling‡",
        "Product labelling; WHO 2025": "Product labelling‡; WHO guideline†",
        "Galli et al, JACC 2025, Supplementary Figure S27, adults without diabetes": "Meta-analysis*, adults without diabetes",
        "Galli et al, JACC 2025, Supplementary Figure S27, adults with diabetes": "Meta-analysis*, adults with diabetes",
        "SELECT glycaemia analysis (Kahn et al, Diabetes Care 2024)": "SELECT glycaemia analysis¶",
        "Network meta-analysis, Clin Obes 2025": "Network meta-analysis‖",
        "Rodriguez et al, JAMA Netw Open 2025": "US cohort of adults using GLP-1RA**",
        "Derived from Rodriguez et al, JAMA Netw Open 2025 (two-year cumulative values)": "Derived from US cohort** (two-year cumulative values)",
        "Weight regain after withdrawal of semaglutide (Wilding et al, Diabetes Obes Metab 2022; Rubino et al, JAMA 2021)": "Weight regain after withdrawal of semaglutide††‡‡",
        "Emerging Risk Factors Collaboration, Lancet 2010": "Collaborative meta-analysis of prospective studies§§",
        "Assumption (scenarios in Figure 5)": "Assumption, varied in scenario analyses",
        "Assumption; correction using time on drug in SELECT in sensitivity analysis (Figure 5)": "Assumption, with a correction for time on drug in SELECT in a sensitivity analysis"}
ws_, rr = _rows(1); col = [c.value for c in ws_[2]].index("Source") + 1; k_ = 0
for i in range(3, ws_.max_row):
    v = ws_.cell(i, col).value
    if v in SRC1: ws_.cell(i, col).value = SRC1[v]; k_ += 1
assert k_ == 21, ("Table 1 source cells converted", k_)
REF1 = ("*Galli M, Benenati S, Laudani C, et al. Cardiovascular effects and tolerability of GLP-1 receptor agonists, a systematic "
        "review and meta-analysis of 99,599 patients. J Am Coll Cardiol 2025;86:1805-19 (supplementary figure S27). "
        "†Celletti F, Farrar J, De Regil L. World Health Organization guideline on the use and indications of glucagon-like peptide-1 "
        "therapies for the treatment of obesity in adults. JAMA 2026;335:434-8. "
        "‡Novo Nordisk. Wegovy (semaglutide) injection, prescribing information. "
        "§Lincoff AM, Brown-Frandsen K, Colhoun HM, et al. Semaglutide and cardiovascular outcomes in obesity without diabetes. "
        "N Engl J Med 2023;389:2221-32. "
        "¶Kahn SE, Deanfield JE, Jeppesen OK, et al. Effect of semaglutide on regression and progression of glycemia in people with "
        "overweight or obesity but without diabetes in the SELECT trial. Diabetes Care 2024;47:1350-9. "
        "‖Ali A, Siddiqui AA, Usman MS, et al. Effect of glucagon-like peptide 1 receptor agonists on systolic blood pressure in "
        "patients with obesity, with or without diabetes, a systematic review and network meta-analysis. Clin Obes 2025. "
        "**Rodriguez PJ, Zhang V, Gratzl S, et al. Discontinuation and reinitiation of dual-labeled GLP-1 receptor agonists among US "
        "adults with overweight or obesity. JAMA Netw Open 2025;8:e2457349. "
        "††Wilding J, Batterham RL, Davies M, et al. Weight regain and cardiometabolic effects after withdrawal of semaglutide, the "
        "STEP 1 trial extension. Diabetes Obes Metab 2022;24:1553-64. "
        "‡‡Rubino D, Abrahamsson N, Davies M, et al. Effect of continued weekly subcutaneous semaglutide vs placebo on weight loss "
        "maintenance in adults with overweight or obesity, the STEP 4 randomized clinical trial. JAMA 2021;325:1414-25. "
        "§§Emerging Risk Factors Collaboration. Diabetes mellitus, fasting blood glucose concentration, and risk of vascular disease, "
        "a collaborative meta-analysis of 102 prospective studies. Lancet 2010;375:2215-22.")
n1 = _note(1); _setnote(1, n1 + " " + REF1)
# Table 2: external benchmarking rather than validation
ws_ = wb["Supplementary Table 2"]; ws_.cell(1, 1).value = str(ws_.cell(1, 1).value).replace("Calibration and validation", "Calibration and external benchmarking")
for i in range(3, ws_.max_row):
    for j in range(1, ws_.max_column + 1):
        v = ws_.cell(i, j).value
        if isinstance(v, str):
            v2 = v.replace("Persistence reconstruction (Rodriguez 2025)", "Persistence reconstruction*").replace("External validation", "External benchmarking").replace("external validation", "external benchmarking")
            ws_.cell(i, j).value = v2
n2 = _note(2).replace("External validation simulated", "External benchmarking simulated")
_setnote(2, n2 + " This comparison checks that the modelled magnitude is plausible and is not a formal external validation, because "
         "the populations and end points are not fully matched. *Rodriguez PJ, Zhang V, Gratzl S, et al. Discontinuation and "
         "reinitiation of dual-labeled GLP-1 receptor agonists among US adults with overweight or obesity. JAMA Netw Open 2025;8:e2457349.")
# Table 3: why dietary multipliers are independent across outcomes and shared across countries
add_note(3, "Dietary multipliers were drawn independently for each outcome because GBD 2023 estimates the dose-response uncertainty "
            "separately for each outcome, and were shared across countries because the same dose-response curves apply to every "
            "country.", at=" CKD, chronic kidney dise")
# Table 14: the not-estimable sentence generated from the values shown
ws_ = wb["Supplementary Table 14"]; ne_ = {}
for i in range(3, ws_.max_row):
    vals = [ws_.cell(i, j).value for j in (4, 5, 6)]
    if "Not estimable" in vals: ne_[ws_.cell(i, 2).value] = [nm for nm, v in zip(("ratio", "lower bound", "upper bound"), vals) if v == "Not estimable"]
n14 = _note(14); old14 = ("For CKD in Bangladesh, Egypt, Nigeria, and South Africa, the upper uncertainty bound is not estimable because the "
                          "high-education PAF approaches zero in some draws; dietary effects on CKD act through sodium only.")
assert old14 in n14, "Table 14 note sentence not found"
desc = " and ".join(f"the {' and '.join(v)} in {c}" for c, v in ne_.items())
n14 = n14.replace(old14, f"For CKD, {desc} are not estimable because the high-education PAF is zero or approaches zero in some draws. "
                  "Dietary effects on CKD act through sodium only.")
_setnote(14, n14)
# Table 16: allocation formula, ISO3 and country names
ws_, rr = _rows(16); hdr16 = [c.value for c in ws_[2]]
H16 = ["Outcome", "Country", "ISO3"] + hdr16[2:]
rows16 = [[r[0], NAME.get(r[1], r[1]), r[1]] + r[2:] for r in rr]
n16 = _note(16)
old16 = "Attributable deaths in each education group were estimated as group PAF x group population share x total deaths from each outcome."
assert old16 in n16, "Table 16 note sentence not found"
n16 = n16.replace(old16, "Attributable deaths in each education group were allocated from the national total in proportion to the group PAF "
                  "multiplied by the group population share, that is, deaths in group g = national attributable deaths x PAF(g) x share(g) / "
                  "sum over groups of PAF x share.")
write_table(ws_, str(ws_.cell(1, 1).value), H16, rows16, n16, widths=[12, 18, 7, 16, 12, 12, 14, 12, 14, 14])
# Table 17: rebuilt from the affordability matrix and the price file (self-pay prices, October 2026)
ws_ = wb["Supplementary Table 17"]; t17_title = str(ws_.cell(1, 1).value)
am = pd.read_csv("data/results/affordability_matrix.csv"); pr = pd.read_csv("data/params/glp1_prices.csv").set_index("iso3")
am = am.assign(_c=am.iso3.map({k: i for i, k in enumerate(ORDER)})).sort_values(["_c", "quintile"])
def _pen(iso):
    r_ = pr.loc[iso]
    return "Not launched" if r_.price_type == "Imputed" else f"{r_.currency} {r_.price_per_pen_lcu:,.2f}"
H17 = ["Country", "ISO3", "Income quintile", "Per-capita income (USD)", "Branded price per pen (local currency)", "Branded annual cost (USD)",
       "Branded price basis", "Branded cost as share of income (%)", "Generic upper bound ($140/yr) as share of income (%)",
       "Generic lower bound ($28/yr) as share of income (%)"]
rows17 = [[NAME.get(r.iso3, r.iso3), r.iso3, r.quintile, round(r.income_pc_usd), _pen(r.iso3), int(r.brand_annual_usd), pr.loc[r.iso3, "price_type"],
           r.afford_brand_pct, r.afford_generic_hi_pct, r.afford_generic_lo_pct] for r in am.itertuples()]
src = "; ".join(f"{NAME.get(i_, i_)}, {pr.loc[i_, 'source_note']}" for i_ in sorted(pr.index, key=lambda k: NAME.get(k, k)))
n17 = ("Per-capita income by quintile was estimated from World Bank income quintile shares (2022-2024) multiplied by GNI per capita "
       "(Atlas method, 2025). Branded prices are those paid by a self-paying patient for semaglutide 2.4 mg, one pen of four weekly doses, "
       "accessed on 27 October 2026, with 13 pens a year, and were converted to US dollars at European Central Bank reference rates of "
       "26 October 2026 (Brazil, 23 October 2026, and Egypt, the approximate official rate). Prices were taken from official price lists "
       "where available and otherwise from pharmacy or clinic retail prices, including value added tax where it applies. In Bangladesh "
       "and Nigeria, where the drug has not been launched, prices were imputed. Generic price bounds (US$140 a year upper, US$28 a year "
       "lower) are from a production-cost analysis (Levi et al, Obesity 2026). Affordability was defined as an annual treatment cost not "
       "exceeding 10% of estimated per-capita income within each income quintile. This study-specific threshold is not a measure of "
       "catastrophic health expenditure under WHO definitions, which relate household out-of-pocket health spending to household "
       "consumption or discretionary budget. Income quintile Q1 is the poorest 20% and Q5 is the richest 20%. IND and NGA quintile incomes "
       "are derived from consumption surveys and may conservatively overestimate affordability for Q1-Q2. Price sources: " + src + ". "
       "GNI, gross national income; GLP-1RA, glucagon-like peptide-1 receptor agonist.")
write_table(ws_, t17_title, H17, rows17, n17, widths=[18, 7, 10, 13, 16, 13, 24, 14, 16, 16])
# references to figures in the notes (tables must be readable on their own)
for n, old, new in [
        (5, "These values give the dietary composition in Figure 1c and the share of the burden targeted by the feasible package (sodium, whole grains and sugar-sweetened beverages) in Figure 4.",
            "These values give the dietary composition of the burden and the share targeted by the feasible package (sodium, whole grains and sugar-sweetened beverages)."),
        (12, "Scenarios are defined in the Methods and shown in Figure 5. ", ""),
        (14, "CKD is excluded from Figure 6a because of unstable educational gradients but is included here for completeness.",
            "Educational gradients for CKD are less stable than for the other outcomes."),
        (21, "Expenditure was calculated as in Supplementary Table 20, treatment paying only", "Expenditure was calculated with treatment paying only"),
        (1, ", following the evidence hierarchy described in the Methods.", "."),
        (1, " (see the table of parameters varied)", ""),
        (1, " (see the calibration and validation table)", "")]:
    t_ = _note(n)
    if old in t_: _setnote(n, t_.replace(old, new))
    else: print(f"[warning] Table {n}: sentence not found: {old[:60]}")

# Table 8: uncertainty intervals of the multiple for T2D and CKD, when available
fo = "data/results/v3/overlap/overlap_ui_t2dckd.csv"
if os.path.exists(fo):
    ou = pd.read_csv(fo); OUT = {"Incident type 2 diabetes cases": "T2D", "Deaths from CKD": "CKD"}
    U = {(r.outcome, r.iso3, r.coverage, int(r.overlap_pct)): (r.mult_lo, r.mult_hi) for r in ou.itertuples()}
    ws_ = wb["Supplementary Table 8"]; h8 = [c.value for c in ws_[2]]; jl, jh = h8.index("Multiple, 95% UI lower") + 1, h8.index("Multiple, 95% UI upper") + 1; k_ = 0
    for i in range(3, ws_.max_row):
        o_ = OUT.get(ws_.cell(i, 3).value)
        if o_:
            lo, hi = U[(o_, ws_.cell(i, 2).value, ws_.cell(i, 4).value, int(ws_.cell(i, 5).value))]
            ws_.cell(i, jl).value, ws_.cell(i, jh).value = _r(lo, 2), _r(hi, 2); k_ += 1
    n8 = _note(8).replace("; uncertainty intervals of the multiple were not estimated for T2D and CKD", ", with multiples computed within each simulation")
    _setnote(8, n8); print(f"Table 8: {k_} T2D and CKD rows given intervals of the multiple")
# Table 20: scope of the expenditure analysis
add_note(20, "Expenditure covers only the drug cost for time on treatment and the implementation cost of the dietary package, from a "
             "public payer perspective, in 2024-25 US dollars, with unit costs held constant over 30 years. It excludes delivery, "
             "monitoring, adverse events and downstream health-care costs and savings, so it is not a full cost-effectiveness analysis.",
         at=" CKD, chronic kidney disease;")
# Table 21: three panels, each with its own header
ws21 = wb["Supplementary Table 21"]
PA = (["Country", "ISO3", "GLP-1RA price", "Dietary policy cost over 30 years (US$ millions)"] + UI3("Deaths averted by dietary policy")
      + ["Treatment coverage bought with the dietary policy cost (% of eligible adults)"] + UI3("Deaths averted by that treatment"))
PB = (["Country", "ISO3", "GLP-1RA price", "Budget, treatment at 20% coverage (US$ millions)", "Dietary policy cost as share of the budget (%)"]
      + UI3("Deaths averted, treatment only") + ["Treatment coverage with the remaining budget (% of eligible adults)"]
      + UI3("Deaths averted, dietary policy plus treatment"))
PC = ["Country", "ISO3", "GLP-1RA price"] + UI3("Deaths averted per dollar, dietary policy relative to treatment")
full = [_row21(r) for r in bq.itertuples()]          # columns as in H21
ix = {h: i for i, h in enumerate(H21)}
def pick(row, cols): return [row[ix[c]] for c in cols]
A_ = ["Country", "ISO3", "GLP-1RA price", "Dietary policy cost over 30 years (US$ millions)", "Deaths averted by dietary policy", "95% UI lower", "95% UI upper",
      "A. Treatment coverage bought with the dietary policy cost (% of eligible adults)", "A. Deaths averted by that treatment"]
rowsA = [row[:4] + row[4:7] + [row[7]] + row[8:11] for row in full]
rowsB = [row[:3] + [row[11], row[12]] + row[13:16] + [row[16]] + row[17:20] for row in full]
rowsC = [row[:3] + row[20:23] for row in full]
leg21 = _note(21)
for rng in list(ws21.merged_cells.ranges): ws21.unmerge_cells(str(rng))
ws21.delete_rows(1, ws21.max_row)
r0 = 1; ws21.cell(r0, 1, T21); _style(ws21.cell(r0, 1), STY["title"]); ws21.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=len(PB)); r0 += 1
for lab, H, R_ in (("A. Treatment funded by the cost of the dietary package", PA, rowsA),
                   ("B. Budget for treatment at 20% coverage, spent on treatment alone or on the dietary package first", PB, rowsB),
                   ("C. Deaths averted per dollar spent, dietary policy relative to treatment at 20% coverage", PC, rowsC)):
    ws21.cell(r0, 1, lab); _style(ws21.cell(r0, 1), STY["title"]); r0 += 1
    for jj, h in enumerate(H, 1): ws21.cell(r0, jj, h); _style(ws21.cell(r0, jj), STY["head"])
    ws21.row_dimensions[r0].height = 60; r0 += 1
    for row in R_:
        for jj, v in enumerate(row, 1): ws21.cell(r0, jj, v); _style(ws21.cell(r0, jj), STY["data"])
        r0 += 1
    r0 += 1
leg21 = leg21.replace("A, the GLP-1RA coverage", "Panel A, the GLP-1RA coverage").replace(" B, the expenditure", " Panel B, the expenditure").replace(" C, deaths averted per dollar", " Panel C, deaths averted per dollar")
ws21.cell(r0, 1, leg21); _style(ws21.cell(r0, 1), STY["note"]); ws21.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=len(PB))
ws21.row_dimensions[r0].height = 150
for jj, w in enumerate([22, 7, 14, 14, 13, 12, 12, 14, 13, 12, 12, 13, 12], 1): ws21.column_dimensions[ws21.cell(2, jj).column_letter].width = w

os.makedirs(os.path.dirname(A.out) or ".", exist_ok=True)
wb.save(A.out)
print(
    "renumbered (old -> new):",
    {o: n for o, n in RENUM.items() if o != n},
    "| dropped:",
    sorted(DROP),
)
print(f"{len(NEW)} tables updated -> {A.out}  ({N_DRAWS} draws)")
for n, a_, b_ in report:
    print(f"  Table {n:2d} (now {RENUM[n]:2d}): rows {a_} -> {b_}")
print(
    f"Full-model calibration: first-year IHD and IS deaths differ from the calibrated rates by a median of "
    f"{CAL['median']:.2f}% (largest {CAL['max']:.2f}%); four-state life table, largest error {err4:.1e}"
)
