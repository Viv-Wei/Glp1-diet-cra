import sys, numpy as np, pandas as pd


def csv_same(a, b):
    x, y = pd.read_csv(a, low_memory=False), pd.read_csv(b, low_memory=False)
    ok, msg = _csv_same(x, y)
    if ok or "rows" in msg:
        return ok, msg
    keys = [c for c in x.columns if c in y.columns and x[c].dtype.kind not in "fiub"]
    if keys:
        xs, ys = x.sort_values(keys, kind="stable").reset_index(
            drop=True
        ), y.sort_values(keys, kind="stable").reset_index(drop=True)
        ok2, msg2 = _csv_same(xs, ys)
        return ok2, msg2 + (
            " (rows in a different order)" if ok2 else " (after sorting rows)"
        )
    return ok, msg


def _csv_same(x, y):
    common = [c for c in x.columns if c in y.columns]
    note = (
        ""
        if list(x.columns) == list(y.columns)
        else f" (compared {len(common)} common columns)"
    )
    if len(x) != len(y):
        return False, f"rows {len(x)} vs {len(y)}"
    for c in common:
        u, v = x[c], y[c]
        if u.dtype.kind in "fi" and v.dtype.kind in "fi":
            if not np.allclose(
                u.values, v.values, rtol=1e-12, atol=1e-12, equal_nan=True
            ):
                return (
                    False,
                    f"column {c}: max |diff| {np.nanmax(np.abs(u.values - v.values)):.3g}",
                )
        elif not u.astype(str).equals(v.astype(str)):
            return False, f"column {c}: text differs"
    return True, "identical" + note


def xlsx_same(a, b):
    from openpyxl import load_workbook

    x, y = load_workbook(a, read_only=True, data_only=True), load_workbook(
        b, read_only=True, data_only=True
    )
    if x.sheetnames != y.sheetnames:
        return False, "sheets differ"
    for s in x.sheetnames:
        if [r for r in x[s].iter_rows(values_only=True)] != [
            r for r in y[s].iter_rows(values_only=True)
        ]:
            return False, f"{s} differs"
    return True, "identical"


bad = 0
for a, b in zip(sys.argv[1::2], sys.argv[2::2]):
    try:
        ok, msg = (xlsx_same if a.endswith(".xlsx") else csv_same)(a, b)
    except FileNotFoundError as e:
        ok, msg = False, f"missing: {e.filename}"
    bad += not ok
    print(f"  {'OK  ' if ok else 'DIFF'} {a}: {msg}")
sys.exit(1 if bad else 0)
